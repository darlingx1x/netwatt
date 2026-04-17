from __future__ import annotations

from collections.abc import AsyncIterator
from io import BytesIO

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook
from PyPDF2 import PdfReader
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from netwatt.auth.models import AuthToken
from netwatt.auth.schemas import RegisterReq
from netwatt.auth.service import register
from netwatt.catalog.models import Equipment
from netwatt.scenarios.calc_runner import perform_calculation
from netwatt.scenarios.models import Scenario
from netwatt.users.models import User


@pytest_asyncio.fixture(autouse=True)
async def _clean(session: AsyncSession) -> AsyncIterator[None]:
    await session.execute(delete(AuthToken))
    await session.execute(delete(Scenario))
    await session.execute(delete(Equipment))
    await session.execute(delete(User))
    await session.commit()
    yield


@pytest_asyncio.fixture()
async def engineer_headers(
    client: AsyncClient, session: AsyncSession
) -> dict[str, str]:
    await register(
        session,
        RegisterReq(email="e@x.com", full_name="E", password="pass1234", role="engineer"),
    )
    await session.commit()
    r = await client.post("/api/auth/login", json={"email": "e@x.com", "password": "pass1234"})
    return {"Authorization": f"Bearer {r.json()['tokens']['access_token']}"}


@pytest_asyncio.fixture()
async def ready_scenario(
    session: AsyncSession,
    client: AsyncClient,
    engineer_headers: dict[str, str],
) -> int:
    eq = Equipment(
        vendor="Cisco", model="Catalyst 9200L", category="access_switch",
        ports_total=24, poe_ports=24, poe_budget_w=370,
        p_idle_w=45, p_max_w=435,
        eee_supported=True, alr_supported=True, poe_scheduling=True,
    )
    session.add(eq)
    await session.commit()
    payload = {
        "name": "Report test",
        "tariff": {"day": 1050, "peak": 1450, "night": 450, "currency": "UZS"},
        "traffic": {
            "day_util": 0.35, "peak_util": 0.7, "night_util": 0.05,
            "day_hours": 6, "peak_hours": 2, "night_hours": 16,
        },
        "policies": {
            "eee": {"enabled": True, "eta": 0.5},
            "alr": {"enabled": True, "drop": 0.4},
            "poe_sched": {"enabled": True, "off_hours": 12},
            "consolidation": {"enabled": False},
        },
        "ef_grid": 0.468,
        "items": [{"equipment_id": eq.id, "quantity": 6}],
    }
    r = await client.post("/api/scenarios", json=payload, headers=engineer_headers)
    sid = r.json()["id"]
    await perform_calculation(session, sid)
    return sid


@pytest.mark.asyncio()
async def test_pdf_generates_and_is_readable(
    client: AsyncClient, engineer_headers: dict[str, str], ready_scenario: int
) -> None:
    r = await client.get(
        f"/api/scenarios/{ready_scenario}/report.pdf?lang=ru", headers=engineer_headers
    )
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/pdf")
    assert len(r.content) > 20_000
    reader = PdfReader(BytesIO(r.content))
    assert len(reader.pages) >= 1


@pytest.mark.asyncio()
async def test_pdf_lang_uz(
    client: AsyncClient, engineer_headers: dict[str, str], ready_scenario: int
) -> None:
    r = await client.get(
        f"/api/scenarios/{ready_scenario}/report.pdf?lang=uz", headers=engineer_headers
    )
    assert r.status_code == 200
    assert len(r.content) > 15_000


@pytest.mark.asyncio()
async def test_xlsx_has_four_sheets(
    client: AsyncClient, engineer_headers: dict[str, str], ready_scenario: int
) -> None:
    r = await client.get(
        f"/api/scenarios/{ready_scenario}/report.xlsx", headers=engineer_headers
    )
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert len(wb.sheetnames) == 4


@pytest.mark.asyncio()
async def test_report_without_result_rejected(
    client: AsyncClient,
    engineer_headers: dict[str, str],
    session: AsyncSession,
) -> None:
    eq = Equipment(
        vendor="X", model="Y", category="router",
        p_idle_w=50, p_max_w=100,
    )
    session.add(eq)
    await session.commit()
    payload = {
        "name": "No result",
        "tariff": {"day": 1000, "peak": 1000, "night": 500, "currency": "UZS"},
        "traffic": {
            "day_util": 0.3, "peak_util": 0.6, "night_util": 0.05,
            "day_hours": 8, "peak_hours": 4, "night_hours": 12,
        },
        "policies": {
            "eee": {"enabled": False},
            "alr": {"enabled": False},
            "poe_sched": {"enabled": False},
            "consolidation": {"enabled": False},
        },
        "items": [{"equipment_id": eq.id, "quantity": 1}],
    }
    r = await client.post("/api/scenarios", json=payload, headers=engineer_headers)
    sid = r.json()["id"]
    pdf = await client.get(
        f"/api/scenarios/{sid}/report.pdf", headers=engineer_headers
    )
    assert pdf.status_code == 400
