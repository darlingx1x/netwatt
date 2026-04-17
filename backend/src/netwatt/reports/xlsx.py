from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from netwatt.reports.i18n import translate
from netwatt.scenarios.models import Scenario


def render_xlsx(scenario: Scenario, lang: str = "ru") -> bytes:
    if scenario.result is None:
        raise ValueError("scenario_has_no_result")
    t = translate(lang)
    wb = Workbook()

    # Sheet 1: Inputs
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = t["inputs_section"][:31]
    _header(ws1, 1, t["param"], t["value"])
    rows = [
        (t["tariff_day"], scenario.tariff.get("day", 0)),
        (t["tariff_peak"], scenario.tariff.get("peak", 0)),
        (t["tariff_night"], scenario.tariff.get("night", 0)),
        (t["day_util"], scenario.traffic.get("day_util", 0)),
        (t["peak_util"], scenario.traffic.get("peak_util", 0)),
        (t["night_util"], scenario.traffic.get("night_util", 0)),
        (t["ef_grid"], float(scenario.ef_grid)),
    ]
    for i, (k, v) in enumerate(rows, start=2):
        ws1.cell(row=i, column=1, value=k)
        ws1.cell(row=i, column=2, value=v)

    # Sheet 2: KPI
    ws2 = wb.create_sheet(t["kpi_section"][:31])
    _header(ws2, 1, t["param"], t["value"])
    r = scenario.result
    kpi = [
        (t["kpi_savings_kwh"], float(r.savings_kwh)),
        (t["kpi_savings_money"], float(r.savings_money)),
        (t["kpi_co2"], float(r.co2_saved_kg)),
        ("E_base, kWh", float(r.e_base_kwh)),
        ("E_opt, kWh", float(r.e_optimized_kwh)),
    ]
    for i, (k, v) in enumerate(kpi, start=2):
        ws2.cell(row=i, column=1, value=k)
        ws2.cell(row=i, column=2, value=v)

    # Sheet 3: Policies breakdown
    ws3 = wb.create_sheet(t["policies_section"][:31])
    _header(ws3, 1, t["policy"], "ΔE kWh/year")
    for i, (k, v) in enumerate(
        [
            ("EEE", r.breakdown.get("eee", 0)),
            ("ALR", r.breakdown.get("alr", 0)),
            ("PoE", r.breakdown.get("poe", 0)),
            ("Consolidation", r.breakdown.get("consolidation", 0)),
        ],
        start=2,
    ):
        ws3.cell(row=i, column=1, value=k)
        ws3.cell(row=i, column=2, value=float(v))

    # Sheet 4: Per device
    ws4 = wb.create_sheet(t["per_device_section"][:31])
    _header(ws4, 1, t["vendor"], t["model"], t["qty"], "E_base", "E_opt", "ΔE")
    for i, d in enumerate(r.per_device, start=2):
        ws4.cell(row=i, column=1, value=d["vendor"])
        ws4.cell(row=i, column=2, value=d["model"])
        ws4.cell(row=i, column=3, value=d["quantity"])
        ws4.cell(row=i, column=4, value=d["e_base_kwh"])
        ws4.cell(row=i, column=5, value=d["e_optimized_kwh"])
        ws4.cell(row=i, column=6, value=d["delta_total_kwh"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _header(ws, row: int, *values: str) -> None:  # type: ignore[no-untyped-def]
    fill = PatternFill("solid", fgColor="1E3A8A")
    font = Font(color="FFFFFF", bold=True)
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(horizontal="center")
